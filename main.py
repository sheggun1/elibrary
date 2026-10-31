import os
import io
import json
import time
import random
import hmac
import html as html_lib
from collections import defaultdict, deque
from typing import List, Optional
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, HTTPException, Header, UploadFile, File, Depends, Request
from fastapi.responses import Response, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from sqlalchemy import (
    create_engine, Column, Integer, String, Text, DateTime,
    ForeignKey, select
)
from sqlalchemy.orm import declarative_base, sessionmaker, Session, relationship
from google import genai
from google.genai import types as genai_types
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_ORIENT
from docx.table import Table
from docx.text.paragraph import Paragraph
from bs4 import BeautifulSoup
from pypdf import PdfReader
import uvicorn

# ---------- Config ----------
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
ADMIN_KEY = os.getenv("ADMIN_KEY", "").strip()
SCHOOL_NAME = os.getenv("SCHOOL_NAME", "").strip()
RATE_LIMIT_PER_HOUR = int(os.getenv("RATE_LIMIT_PER_HOUR", "40"))
if not GEMINI_API_KEY:
    raise RuntimeError("GEMINI_API_KEY environment variable is not set")
if not ADMIN_KEY:
    print("WARNING: ADMIN_KEY is not set. Admin sign-in is disabled until you set it.")

client = genai.Client(api_key=GEMINI_API_KEY)
MODEL_NAME = os.getenv("MODEL_NAME", "gemini-3.8-flash")

DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./curriculum.db")
# Render (and Heroku) hand out postgres:// URLs; SQLAlchemy needs postgresql://
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"

# ---------- Database ----------
engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True,
    pool_recycle=300,  # serverless Postgres (Neon) closes idle connections
    connect_args={"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
Base = declarative_base()

LIST_FIELDS = [
    "performance_objectives", "content_knowledge", "key_competencies",
    "learning_activities", "teaching_resources", "evaluation_guide",
]


class Curriculum(Base):
    __tablename__ = "curricula"
    id = Column(Integer, primary_key=True, index=True)
    class_level = Column(String(50), nullable=False)
    subject = Column(String(255), nullable=False)
    term = Column(String(50), nullable=False, default="First Term")
    created_at = Column(DateTime, default=datetime.utcnow)
    weeks = relationship("Week", back_populates="curriculum", cascade="all, delete-orphan")


class Week(Base):
    __tablename__ = "weeks"
    id = Column(Integer, primary_key=True, index=True)
    curriculum_id = Column(Integer, ForeignKey("curricula.id"), nullable=False)
    week_number = Column(Integer, nullable=False)
    topic = Column(String(500), nullable=False)
    performance_objectives = Column(Text, default="[]")
    content_knowledge = Column(Text, default="[]")
    key_competencies = Column(Text, default="[]")
    learning_activities = Column(Text, default="[]")
    teaching_resources = Column(Text, default="[]")
    evaluation_guide = Column(Text, default="[]")
    curriculum = relationship("Curriculum", back_populates="weeks")


DB_STARTUP_ERROR = None
try:
    Base.metadata.create_all(bind=engine)
except Exception as _e:  # keep the site up so /api/status can say what is wrong
    DB_STARTUP_ERROR = f"{type(_e).__name__}: {str(_e).splitlines()[0][:200]}"
    print("DATABASE ERROR:", DB_STARTUP_ERROR)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def week_to_dict(w: Week):
    return {
        "id": w.id,
        "curriculum_id": w.curriculum_id,
        "week_number": w.week_number,
        "topic": w.topic,
        **{f: json.loads(getattr(w, f) or "[]") for f in LIST_FIELDS},
    }


def curriculum_to_dict(c: Curriculum, with_weeks: bool = True):
    d = {
        "id": c.id,
        "class_level": c.class_level,
        "subject": c.subject,
        "term": c.term,
        "weeks_count": len(c.weeks),
        "created_at": c.created_at.isoformat() if c.created_at else None,
    }
    if with_weeks:
        d["weeks"] = [week_to_dict(w) for w in sorted(c.weeks, key=lambda x: x.week_number)]
    return d


# ---------- Pydantic models ----------
class WeekInput(BaseModel):
    week_number: int
    topic: str
    performance_objectives: List[str] = []
    content_knowledge: List[str] = []
    key_competencies: List[str] = []
    learning_activities: List[str] = []
    teaching_resources: List[str] = []
    evaluation_guide: List[str] = []


class CurriculumInput(BaseModel):
    class_level: str
    subject: str
    term: str = "First Term"
    weeks: List[WeekInput]


class LessonNoteRequest(BaseModel):
    class_level: str
    subject: str
    term: str
    week_number: int
    topic: str
    performance_objectives: List[str]
    content_knowledge: List[str]
    key_competencies: List[str]
    learning_activities: List[str]
    teaching_resources: List[str]
    evaluation_guide: List[str]


class QuestionRequest(BaseModel):
    week_ids: List[int] = Field(min_length=1, max_length=20)
    objective_count: int = Field(default=10, ge=0, le=50)
    theory_count: int = Field(default=3, ge=0, le=15)
    difficulty: str = "mixed"  # easy | medium | hard | mixed


class CustomTopicRequest(BaseModel):
    class_level: str = Field(min_length=1, max_length=50)
    subject: str = Field(min_length=1, max_length=255)
    topic: str = Field(min_length=2, max_length=500)
    details: str = Field(default="", max_length=2000)  # optional: sub-topics / objectives the teacher wants covered


class CustomQuestionRequest(CustomTopicRequest):
    objective_count: int = Field(default=10, ge=0, le=50)
    theory_count: int = Field(default=3, ge=0, le=15)
    difficulty: str = "mixed"


class ExportRequest(BaseModel):
    html: str
    title: str = "Document"
    landscape: bool = False


# ---------- Auth / rate limiting ----------
def _is_admin_key(key: Optional[str]) -> bool:
    return bool(ADMIN_KEY) and bool(key) and hmac.compare_digest(key.strip().encode(), ADMIN_KEY.encode())


def require_admin(x_admin_key: Optional[str] = Header(None)):
    if not ADMIN_KEY:
        raise HTTPException(
            status_code=503,
            detail="ADMIN_KEY is not set on the server. Add it in Render > Environment, then redeploy.",
        )
    if not _is_admin_key(x_admin_key):
        raise HTTPException(status_code=401, detail="Invalid or missing X-Admin-Key header")
    return True


_hits = defaultdict(deque)


def rate_limit(request: Request):
    """Protects the (paid) AI endpoints from abuse. Admins are exempt."""
    if _is_admin_key(request.headers.get("x-admin-key")):
        return
    fwd = request.headers.get("x-forwarded-for", "").split(",")[0].strip()
    ip = fwd or (request.client.host if request.client else "unknown")
    now = time.time()
    q = _hits[ip]
    while q and now - q[0] > 3600:
        q.popleft()
    if len(q) >= RATE_LIMIT_PER_HOUR:
        raise HTTPException(status_code=429, detail="Too many generation requests. Please try again later.")
    q.append(now)


# ---------- Gemini helpers ----------
def _strip_fences(text: str) -> str:
    t = (text or "").strip()
    if t.startswith("```"):
        t = t.split("\n", 1)[1] if "\n" in t else t[3:]
    if t.endswith("```"):
        t = t[:-3]
    return t.strip()


def _ai_hint(err: str) -> str:
    low = err.lower()
    hint = ""
    if "api key not valid" in low or "api_key_invalid" in low or "permission_denied" in low or "403" in err:
        hint = " -> Check GEMINI_API_KEY in Render > Environment (no spaces, key still active)."
    elif "404" in err or "not found" in low:
        hint = f" -> The model '{MODEL_NAME}' was not found for this key. Set a MODEL_NAME environment variable to a model your key can use."
    elif "429" in err or "quota" in low or "resource_exhausted" in low:
        hint = " -> Gemini quota or rate limit reached. Wait a minute or check your Google AI Studio quota."
    return err[:400] + hint


def call_gemini(prompt: str, *, max_tokens: int, temperature: float, json_mode: bool = False) -> str:
    cfg = dict(max_output_tokens=max_tokens, temperature=temperature)
    if json_mode:
        cfg["response_mime_type"] = "application/json"
    max_retries = 4
    base_delay = 1.0
    for attempt in range(max_retries + 1):
        try:
            response = client.models.generate_content(
                model=MODEL_NAME,
                contents=prompt,
                config=genai_types.GenerateContentConfig(**cfg),
            )
            if not response.text:
                raise RuntimeError("The AI returned an empty answer (it may have been blocked or cut off)")
            return _strip_fences(response.text)
        except Exception as e:
            error_str = str(e)
            is_retryable = any(k in error_str for k in ["503", "429", "UNAVAILABLE", "high demand"])
            if is_retryable and attempt < max_retries:
                delay = base_delay * (2 ** attempt) + random.uniform(0, 1)
                print(f"Attempt {attempt + 1} failed. Retrying in {delay:.1f}s...")
                time.sleep(delay)
                continue
            raise HTTPException(status_code=500, detail=f"AI generation failed: {_ai_hint(error_str)}")
    raise HTTPException(status_code=500, detail="AI generation failed: max retries exceeded")


def call_gemini_json(prompt: str, *, max_tokens: int, temperature: float) -> dict:
    last_err = None
    for _ in range(2):  # one retry if the model returns malformed JSON
        raw = call_gemini(prompt, max_tokens=max_tokens, temperature=temperature, json_mode=True)
        try:
            return json.loads(raw)
        except json.JSONDecodeError as e:
            last_err = (e, raw)
    e, raw = last_err
    raise HTTPException(status_code=502, detail=f"AI returned invalid JSON: {e}. First 300 chars: {raw[:300]}")


# ---------- HTML sanitising (AI output is rendered in the browser and in PDFs) ----------
ALLOWED_TAGS = {
    "h1", "h2", "h3", "h4", "h5", "p", "ul", "ol", "li", "strong", "b", "em", "i", "u",
    "br", "table", "thead", "tbody", "tr", "th", "td", "div", "span", "sup", "sub", "hr",
}
ALLOWED_ATTRS = {"td": {"colspan", "rowspan"}, "th": {"colspan", "rowspan"}}


def sanitize_html(html_str: str) -> str:
    soup = BeautifulSoup(html_str or "", "html.parser")
    for t in soup.find_all(["script", "style", "iframe", "object", "embed", "link", "meta", "form"]):
        t.decompose()
    for t in soup.find_all(True):
        if t.name not in ALLOWED_TAGS:
            t.unwrap()
            continue
        keep = ALLOWED_ATTRS.get(t.name, set())
        t.attrs = {k: v for k, v in t.attrs.items() if k in keep}
    return str(soup)


# ---------- Lesson Note Prompt ----------
PROMPT_TEMPLATE = """
You are an expert Nigerian teacher. Generate a complete NERDC-compliant LESSON NOTE.

=== CURRICULUM DATA (USE ONLY THIS) ===
Class: {class_level}
Subject: {subject}
Term: {term}
Week: {week_number}
Topic: {topic}
Performance Objectives: {objectives}
Content Knowledge: {content}
Key Competencies: {competencies}
Learning Activities: {activities}
Teaching Resources: {resources}
Evaluation Guide: {evaluation}
=========================================

Output a complete lesson note in HTML with these sections:

1. GENERAL INFORMATION
   Class: {class_level}
   Subject: {subject}
   Term: {term}
   Week: {week_number}
   Topic: {topic}
   Duration: 40 minutes

2. PERFORMANCE OBJECTIVES
   "By the end of the lesson, students should be able to:" followed by a numbered list.

3. INSTRUCTIONAL MATERIALS
   A bulleted list.

4. PREVIOUS KNOWLEDGE
   1-2 sentences in Nigerian context.

5. INTRODUCTION (5 minutes)
   3-4 sentence hook in Nigerian context.

6. LESSON DEVELOPMENT (25 minutes)
   Use this format for EACH step (at least 3 steps):
   <h4>Step N: [Title] (X min)</h4>
   <p><strong>Teacher Activity:</strong> [description]</p>
   <p><strong>Student Activity:</strong> [description]</p>
   DO NOT use HTML tables for this section.

7. ENRICHMENT SECTION
   - 3 Nigerian context examples (as <h4>Nigerian Context Examples:</h4> then bulleted list)
   - 1 group activity (as <h4>Group Activity:</h4> then paragraph)
   - 1 differentiation tip (as <h4>Differentiation Tip:</h4> then paragraph)

8. EVALUATION (5 minutes)
   Numbered list of 3-5 verbal questions from the Evaluation Guide.

9. SUMMARY & CONCLUSION
   2-3 sentences.

10. ASSIGNMENT
    1 practical homework task.

If Term is blank or Week is 0, leave those lines out of GENERAL INFORMATION.
Use simple English appropriate to {class_level}. Do not invent topics outside the data.
Write mathematical expressions in plain text/unicode (e.g. x², √25, 3/4), not LaTeX.
Return ONLY the HTML body content (no <html> or <body> tags, just inner HTML).
"""


def generate_lesson_html(req: LessonNoteRequest) -> str:
    def j(items):
        return ", ".join(items) if items else "(none given - work out suitable ones from the topic)"
    prompt = PROMPT_TEMPLATE.format(
        class_level=req.class_level,
        subject=req.subject,
        term=req.term,
        week_number=req.week_number,
        topic=req.topic,
        objectives=j(req.performance_objectives),
        content=j(req.content_knowledge),
        competencies=j(req.key_competencies),
        activities=j(req.learning_activities),
        resources=j(req.teaching_resources),
        evaluation=j(req.evaluation_guide)
    )
    return sanitize_html(call_gemini(prompt, max_tokens=16384, temperature=0.7))


# ---------- Scheme of work (built straight from the curriculum, no AI needed) ----------
def _cell_list(items: List[str]) -> str:
    items = [i for i in items if i and i.strip()]
    if not items:
        return ""
    return "<ul>" + "".join(f"<li>{html_lib.escape(i)}</li>" for i in items) + "</ul>"


def build_scheme_html(cur: Curriculum) -> str:
    e = html_lib.escape
    weeks = sorted(cur.weeks, key=lambda w: w.week_number)
    head = ""
    if SCHOOL_NAME:
        head += f"<h1 style='text-align:center'>{e(SCHOOL_NAME)}</h1>"
    head += (
        f"<h2>SCHEME OF WORK</h2>"
        f"<p><strong>Class:</strong> {e(cur.class_level)} &nbsp;|&nbsp; "
        f"<strong>Subject:</strong> {e(cur.subject)} &nbsp;|&nbsp; "
        f"<strong>Term:</strong> {e(cur.term)}</p>"
    )
    rows = []
    for w in weeks:
        d = week_to_dict(w)
        rows.append(
            "<tr>"
            f"<td>{d['week_number']}</td>"
            f"<td><strong>{e(d['topic'])}</strong></td>"
            f"<td>{_cell_list(d['performance_objectives'])}</td>"
            f"<td>{_cell_list(d['content_knowledge'])}</td>"
            f"<td>{_cell_list(d['learning_activities'])}</td>"
            f"<td>{_cell_list(d['teaching_resources'])}</td>"
            f"<td>{_cell_list(d['evaluation_guide'])}</td>"
            "</tr>"
        )
    table = (
        "<table><thead><tr>"
        "<th>Week</th><th>Topic</th><th>Performance Objectives</th><th>Content</th>"
        "<th>Learning Activities</th><th>Resources</th><th>Evaluation</th>"
        "</tr></thead><tbody>" + "".join(rows) + "</tbody></table>"
    )
    return head + table


# ---------- Question generation ----------
QUESTION_PROMPT = """
You are an experienced Nigerian examiner setting questions for {class_level} {subject} {term}.

=== CURRICULUM CONTENT (SET QUESTIONS ONLY FROM THIS) ===
{curriculum_block}
=========================================================

Generate exactly {n_obj} multiple-choice (objective) questions and exactly {n_theory} theory questions.
Difficulty: {difficulty}. Spread the questions evenly across the topics above.
Use simple, clear English suitable for {class_level}. Use Nigerian context where natural.
Write mathematical expressions in plain text/unicode (e.g. x², √25, 3/4), not LaTeX.

Rules for objective questions:
- Exactly 4 options (A, B, C, D), exactly one correct, plausible distractors.
- Avoid "all of the above" / "none of the above".
- Spread the correct answers across A-D.
Rules for theory questions:
- Give a mark allocation (marks) and a short marking guide listing the key points.

Return ONLY valid JSON, no other text, in exactly this shape:
{{
  "objective": [
    {{"question": "...", "options": {{"A": "...", "B": "...", "C": "...", "D": "..."}}, "answer": "A", "explanation": "..."}}
  ],
  "theory": [
    {{"question": "...", "marks": 5, "marking_guide": ["...", "..."]}}
  ]
}}
"""


def _clean_questions(data: dict, n_obj: int, n_theory: int) -> dict:
    objective = []
    for q in (data.get("objective") or []):
        try:
            opts = q.get("options") or {}
            opts = {k: str(opts[k]).strip() for k in ("A", "B", "C", "D")}
            ans = str(q.get("answer", "")).strip().upper()[:1]
            text = str(q.get("question", "")).strip()
            if text and ans in opts:
                objective.append({
                    "question": text, "options": opts, "answer": ans,
                    "explanation": str(q.get("explanation", "")).strip(),
                })
        except (KeyError, AttributeError, TypeError):
            continue
    theory = []
    for q in (data.get("theory") or []):
        try:
            text = str(q.get("question", "")).strip()
            if not text:
                continue
            try:
                marks = int(q.get("marks", 5))
            except (TypeError, ValueError):
                marks = 5
            guide = q.get("marking_guide") or []
            if isinstance(guide, str):
                guide = [guide]
            theory.append({"question": text, "marks": marks, "marking_guide": [str(g) for g in guide]})
        except AttributeError:
            continue
    return {"objective": objective[:n_obj], "theory": theory[:n_theory]}


def render_questions_html(meta: dict, qs: dict, with_answers: bool) -> str:
    e = html_lib.escape
    out = []
    if SCHOOL_NAME:
        out.append(f"<h1 style='text-align:center'>{e(SCHOOL_NAME)}</h1>")
    out.append(
        f"<h2>{e(meta['subject'])} — {e(meta['class_level'])}{' (' + e(meta['term']) + ')' if meta['term'] else ''}"
        f"{' — MARKING SCHEME' if with_answers else ''}</h2>"
    )
    out.append(f"<p><strong>Topics:</strong> {e(meta['topics'])}</p>")
    if not with_answers:
        out.append("<p>Name: ____________________________ &nbsp; Date: ______________</p>")
    if qs["objective"]:
        out.append("<h3>SECTION A: OBJECTIVE</h3>")
        if not with_answers:
            out.append("<p><em>Choose the correct option.</em></p>")
        out.append("<ol>")
        for q in qs["objective"]:
            opts = "".join(
                f"<li>{k}. {e(v)}</li>" for k, v in q["options"].items()
            )
            ans = ""
            if with_answers:
                ans = f"<p><strong>Answer: {q['answer']}</strong>"
                if q["explanation"]:
                    ans += f" — {e(q['explanation'])}"
                ans += "</p>"
            out.append(f"<li><p>{e(q['question'])}</p><ul>{opts}</ul>{ans}</li>")
        out.append("</ol>")
    if qs["theory"]:
        total = sum(q["marks"] for q in qs["theory"])
        out.append(f"<h3>SECTION B: THEORY ({total} marks)</h3>")
        if not with_answers:
            out.append("<p><em>Answer all questions.</em></p>")
        out.append("<ol>")
        for q in qs["theory"]:
            guide = ""
            if with_answers and q["marking_guide"]:
                guide = "<p><strong>Marking guide:</strong></p><ul>" + "".join(
                    f"<li>{e(g)}</li>" for g in q["marking_guide"]) + "</ul>"
            out.append(f"<li><p>{e(q['question'])} <strong>({q['marks']} marks)</strong></p>{guide}</li>")
        out.append("</ol>")
    return "".join(out)


# ---------- HTML -> DOCX ----------
def html_to_docx(html_str: str, title: str, landscape: bool = False) -> bytes:
    doc = Document()
    if landscape:
        s = doc.sections[0]
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = s.page_height, s.page_width
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, side, Inches(0.6))
    doc.add_heading(title, level=0)
    soup = BeautifulSoup(html_str, "html.parser")
    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "table"]):
        if element.find_parent("table") is not None:
            continue
        if element.name in ("h1", "h2", "h3", "h4"):
            doc.add_heading(element.get_text(strip=True), level=int(element.name[1]))
        elif element.name == "p":
            if element.find_parent("li") is not None:
                continue  # text is emitted with its list item
            text = element.get_text(strip=True)
            if text:
                doc.add_paragraph(text)
        elif element.name == "li":
            parent = element.find_parent(["ol", "ul"])
            # for items that contain nested lists only take their own text
            own = " ".join(
                (c.get_text(" ", strip=True) if hasattr(c, "get_text") else str(c).strip())
                for c in element.children
                if not (getattr(c, "name", None) in ("ul", "ol"))
            ).strip()
            if not own:
                continue
            if parent is not None and parent.name == "ol":
                idx = [li for li in parent.find_all("li", recursive=False)].index(element) + 1
                p = doc.add_paragraph(f"{idx}. {own}")
                p.paragraph_format.left_indent = Inches(0.25)
            else:
                nested = element.find_parent("li") is not None
                p = doc.add_paragraph(own, style="List Bullet")
                if nested:
                    p.paragraph_format.left_indent = Inches(0.75)
        elif element.name == "table":
            rows = element.find_all("tr")
            if rows:
                cols = max(len(r.find_all(["th", "td"])) for r in rows)
                table = doc.add_table(rows=0, cols=cols)
                table.style = "Table Grid"
                for row in rows:
                    cells = row.find_all(["th", "td"])
                    row_cells = table.add_row().cells
                    for i, cell in enumerate(cells):
                        if i < cols:
                            lis = cell.find_all("li")
                            if lis:
                                txt = "\n".join("• " + li.get_text(" ", strip=True) for li in lis)
                            else:
                                txt = cell.get_text(" ", strip=True)
                            row_cells[i].text = txt
                            if cell.name == "th":
                                for r in row_cells[i].paragraphs[0].runs:
                                    r.bold = True
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------- HTML -> PDF ----------
def _deny_fetcher(url, *args, **kwargs):
    raise ValueError("External resources are disabled")


def html_to_pdf(html_str: str, title: str, landscape: bool = False) -> bytes:
    full_html = f"""
    <!DOCTYPE html>
    <html>
    <head><meta charset="UTF-8"><title>{html_lib.escape(title)}</title>
    <style>
        @page {{ size: {'A4 landscape' if landscape else 'A4'}; margin: 15mm; }}
        body {{ font-family: Arial, sans-serif; line-height: 1.5; color: #222; font-size: {'10pt' if landscape else '11pt'}; }}
        h1, h2, h3, h4 {{ color: #1a5276; }}
        h2 {{ border-bottom: 1px solid #ccc; padding-bottom: 4px; margin-top: 18px; }}
        h4 {{ margin-top: 12px; margin-bottom: 4px; }}
        table {{ border-collapse: collapse; width: 100%; margin: 10px 0; page-break-inside: auto; }}
        tr {{ page-break-inside: avoid; }}
        th, td {{ border: 1px solid #333; padding: 6px; text-align: left; vertical-align: top; word-wrap: break-word; }}
        th {{ background: #eaf2f8; }}
        td ul {{ margin: 0; padding-left: 14px; }}
        ul, ol {{ margin: 6px 0 6px 18px; }}
        p {{ margin: 6px 0; }}
    </style>
    </head>
    <body>{sanitize_html(html_str)}</body>
    </html>
    """
    try:
        from weasyprint import HTML  # imported lazily so the app still boots if system libs are missing
        return HTML(string=full_html, url_fetcher=_deny_fetcher).write_pdf()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}")


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _safe_filename(name: str, ext: str) -> str:
    keep = "".join(c if (c.isalnum() or c in "-_ ") else "_" for c in name).strip().replace(" ", "_")
    return f"{keep or 'document'}.{ext}"


# ---------- FastAPI app ----------
app = FastAPI(title="Lesson Note AI Backend")


@app.on_event("startup")
def seed_if_empty():
    if DB_STARTUP_ERROR:
        return
    db = SessionLocal()
    try:
        count = db.execute(select(Curriculum)).scalars().first()
        if count is None:
            cur = Curriculum(
                class_level="JSS 1",
                subject="Solar PV Installation & Maintenance",
                term="First Term"
            )
            db.add(cur)
            db.flush()
            w = Week(
                curriculum_id=cur.id,
                week_number=1,
                topic="Solar Photovoltaics (PV) System",
                performance_objectives=json.dumps([
                    "define Solar Photovoltaics",
                    "identify types of Solar Photovoltaic Modules",
                    "explain how a Solar Photovoltaic system generates electricity"
                ]),
                content_knowledge=json.dumps([
                    "Meaning of Solar Photovoltaics (PV)",
                    "Types of Solar Photovoltaics (PV) Modules",
                    "Solar PV and Electricity generation"
                ]),
                key_competencies=json.dumps(["Communication", "ICT and Digital Competencies"]),
                learning_activities=json.dumps([
                    "Discuss in groups the meaning of Solar PV system",
                    "Watch short videos on types of PV modules",
                    "Demonstrate how PV generates electricity"
                ]),
                teaching_resources=json.dumps([
                    "Internet enabled digital tools",
                    "Sample/Model of Solar PV system components",
                    "Solar Panels, Inverters, Batteries"
                ]),
                evaluation_guide=json.dumps([
                    "explain how a solar PV system generates electricity",
                    "list 2 types of Solar PV Modules"
                ])
            )
            db.add(w)
            db.commit()
            print("Seeded Week 1 of Solar PV JSS 1.")
    finally:
        db.close()


@app.get("/")
def root(request: Request):
    # Browsers get the website; API clients / health pings keep the original JSON.
    if "text/html" in request.headers.get("accept", "") and (STATIC_DIR / "index.html").exists():
        return FileResponse(STATIC_DIR / "index.html")
    return {"status": "alive", "service": "Lesson Note AI", "model": MODEL_NAME, "version": "step4"}


@app.get("/health")
def health():
    return {"status": "healthy", "model": MODEL_NAME}


@app.get("/api/status")
def api_status():
    """Public, secret-free health report so setup problems are visible on the site."""
    db_ok, db_err, count = True, DB_STARTUP_ERROR, None
    if DB_STARTUP_ERROR:
        db_ok = False
    else:
        try:
            with SessionLocal() as db:
                count = len(db.execute(select(Curriculum.id)).scalars().all())
        except Exception as e:
            db_ok, db_err = False, f"{type(e).__name__}: {str(e).splitlines()[0][:200]}"
    return {
        "database_ok": db_ok,
        "database_type": engine.dialect.name,
        "database_error": db_err,
        "curricula_count": count,
        "admin_key_set": bool(ADMIN_KEY),
        "model": MODEL_NAME,
    }


@app.post("/admin/diagnose")
def admin_diagnose(_: bool = Depends(require_admin)):
    """Makes one tiny AI call so the admin can see the real error, if any."""
    try:
        text = call_gemini("Reply with the single word OK.", max_tokens=64, temperature=0)
        return {"ok": True, "model": MODEL_NAME, "reply": text[:50]}
    except HTTPException as e:
        return {"ok": False, "model": MODEL_NAME, "error": e.detail}


# ---------- Admin: curriculum management ----------
def _apply_weeks(db: Session, cur: Curriculum, weeks: List[WeekInput]):
    """Upsert weeks by week_number; weeks missing from the payload are removed."""
    existing = {w.week_number: w for w in cur.weeks}
    seen = set()
    for w in weeks:
        seen.add(w.week_number)
        row = existing.get(w.week_number)
        if row is None:
            row = Week(curriculum_id=cur.id, week_number=w.week_number, topic=w.topic)
            db.add(row)
        row.topic = w.topic
        for f in LIST_FIELDS:
            setattr(row, f, json.dumps([s for s in getattr(w, f) if s and s.strip()]))
    for num, row in existing.items():
        if num not in seen:
            db.delete(row)


@app.post("/admin/curriculum/json")
def admin_save_json(payload: CurriculumInput, _: bool = Depends(require_admin), db: Session = Depends(get_db)):
    class_level = payload.class_level.strip()
    subject = payload.subject.strip()
    term = (payload.term or "First Term").strip()
    if not class_level or not subject:
        raise HTTPException(status_code=400, detail="class_level and subject are required")
    if len({w.week_number for w in payload.weeks}) != len(payload.weeks):
        raise HTTPException(status_code=400, detail="Duplicate week numbers in payload")

    cur = db.execute(
        select(Curriculum).where(
            Curriculum.class_level == class_level,
            Curriculum.subject == subject,
            Curriculum.term == term,
        )
    ).scalars().first()
    replaced = cur is not None
    if cur is None:
        cur = Curriculum(class_level=class_level, subject=subject, term=term)
        db.add(cur)
        db.flush()
    _apply_weeks(db, cur, payload.weeks)
    db.commit()
    return {"ok": True, "curriculum_id": cur.id, "weeks_saved": len(payload.weeks), "replaced": replaced}


@app.get("/admin/curriculum/{curriculum_id}")
def admin_get(curriculum_id: int, _: bool = Depends(require_admin), db: Session = Depends(get_db)):
    cur = db.get(Curriculum, curriculum_id)
    if not cur:
        raise HTTPException(status_code=404, detail="Curriculum not found")
    return curriculum_to_dict(cur)


@app.put("/admin/curriculum/{curriculum_id}")
def admin_update(curriculum_id: int, payload: CurriculumInput,
                 _: bool = Depends(require_admin), db: Session = Depends(get_db)):
    cur = db.get(Curriculum, curriculum_id)
    if not cur:
        raise HTTPException(status_code=404, detail="Curriculum not found")
    if len({w.week_number for w in payload.weeks}) != len(payload.weeks):
        raise HTTPException(status_code=400, detail="Duplicate week numbers in payload")
    cur.class_level = payload.class_level.strip()
    cur.subject = payload.subject.strip()
    cur.term = (payload.term or "First Term").strip()
    _apply_weeks(db, cur, payload.weeks)
    db.commit()
    return {"ok": True, "curriculum_id": cur.id, "weeks_saved": len(payload.weeks)}


EXTRACTION_PROMPT = """
You are a curriculum parser for Nigerian NERDC documents.

Extract ALL weekly curriculum entries from the source text below into STRICT JSON.

Return ONLY valid JSON matching this schema, no other text:
{{
  "class_level": "e.g. JSS 1 / SSS 1 / SSS 2",
  "subject": "e.g. Mathematics / English Language",
  "term": "First Term",
  "weeks": [
    {{
      "week_number": 1,
      "topic": "...",
      "performance_objectives": ["...", "..."],
      "content_knowledge": ["...", "..."],
      "key_competencies": ["...", "..."],
      "learning_activities": ["...", "..."],
      "teaching_resources": ["...", "..."],
      "evaluation_guide": ["...", "..."]
    }}
  ]
}}

CRITICAL RULES:
- Extract EVERY week/topic in the document. Do not stop early.
- If the document contains 12 weeks, return 12 weeks. If 19, return 19.
- If a field is missing for a week, use []. Do NOT leave a sentence or list item unfinished.
- Do NOT truncate. Keep going until the entire document has been processed.
- Preserve the Nigerian curriculum wording where possible.
- If term is not stated, default to "First Term".
- Output pure JSON only (no markdown fences, no explanations).

=== SOURCE TEXT ===
{source_text}
"""


def _docx_text(raw: bytes) -> str:
    doc = Document(io.BytesIO(raw))
    lines = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            t = Paragraph(child, doc).text.strip()
            if t:
                lines.append(t)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                seen, cells = set(), []
                for c in row.cells:
                    if id(c._tc) in seen:  # merged cells repeat
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip().replace("\n", "; "))
                lines.append(" | ".join(cells))
    return "\n".join(lines)


async def _extract_curriculum(file: UploadFile) -> dict:
    name = (file.filename or "").lower()
    raw = await file.read()
    if len(raw) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 15 MB)")
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif name.endswith(".docx"):
            text = _docx_text(raw)
        else:
            raise HTTPException(status_code=400, detail="Only .pdf or .docx files accepted")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read file: {e}")

    if len(text.strip()) < 200:
        raise HTTPException(status_code=400, detail="Extracted text too short — is this a scanned PDF?")

    parsed = call_gemini_json(
        EXTRACTION_PROMPT.format(source_text=text[:200000]),
        max_tokens=32768, temperature=0.1,
    )
    weeks = parsed.get("weeks", [])
    return {
        "ok": True,
        "weeks_extracted": len(weeks),
        "message": "Review and correct the extraction, then save it.",
        "extracted": parsed,
    }


@app.post("/admin/curriculum/pdf")
async def admin_extract_pdf(file: UploadFile = File(...), _: bool = Depends(require_admin)):
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Only .pdf files accepted")
    return await _extract_curriculum(file)


@app.post("/admin/curriculum/file")
async def admin_extract_file(file: UploadFile = File(...), _: bool = Depends(require_admin)):
    return await _extract_curriculum(file)


@app.get("/admin/curricula")
def admin_list(_: bool = Depends(require_admin), db: Session = Depends(get_db)):
    items = db.execute(select(Curriculum).order_by(Curriculum.id)).scalars().all()
    return [curriculum_to_dict(c, with_weeks=False) for c in items]


@app.delete("/admin/curriculum/{curriculum_id}")
def admin_delete(curriculum_id: int, _: bool = Depends(require_admin), db: Session = Depends(get_db)):
    cur = db.get(Curriculum, curriculum_id)
    if not cur:
        raise HTTPException(status_code=404, detail="Curriculum not found")
    db.delete(cur)
    db.commit()
    return {"ok": True, "deleted": curriculum_id}


# ---------- Public read endpoints ----------
@app.get("/curriculum/classes")
def list_classes(db: Session = Depends(get_db)):
    rows = db.execute(select(Curriculum.class_level).distinct()).scalars().all()
    return sorted(rows)


@app.get("/curriculum/{class_level}/subjects")
def list_subjects(class_level: str, db: Session = Depends(get_db)):
    rows = db.execute(
        select(Curriculum.subject).where(Curriculum.class_level == class_level).distinct()
    ).scalars().all()
    return sorted(rows)


@app.get("/curriculum/{class_level}/{subject}/terms")
def list_terms(class_level: str, subject: str, db: Session = Depends(get_db)):
    rows = db.execute(
        select(Curriculum.term)
        .where(Curriculum.class_level == class_level, Curriculum.subject == subject)
        .distinct()
    ).scalars().all()
    return sorted(rows)


@app.get("/curriculum/{class_level}/{subject}/{term}/weeks")
def list_weeks(class_level: str, subject: str, term: str, db: Session = Depends(get_db)):
    cur = db.execute(
        select(Curriculum)
        .where(
            Curriculum.class_level == class_level,
            Curriculum.subject == subject,
            Curriculum.term == term
        )
    ).scalars().first()
    if not cur:
        raise HTTPException(status_code=404, detail="Curriculum not found")
    weeks = sorted(cur.weeks, key=lambda w: w.week_number)
    return [
        {"id": w.id, "week_number": w.week_number, "topic": w.topic}
        for w in weeks
    ]


@app.get("/curriculum/week/{week_id}")
def get_week(week_id: int, db: Session = Depends(get_db)):
    w = db.get(Week, week_id)
    if not w:
        raise HTTPException(status_code=404, detail="Week not found")
    return week_to_dict(w)


# ---------- Website API (used by static/index.html) ----------
@app.get("/api/curricula")
def api_curricula(db: Session = Depends(get_db)):
    items = db.execute(
        select(Curriculum).order_by(Curriculum.class_level, Curriculum.subject, Curriculum.term)
    ).scalars().all()
    return [curriculum_to_dict(c, with_weeks=False) for c in items]


@app.get("/api/curricula/{curriculum_id}/weeks")
def api_curriculum_weeks(curriculum_id: int, db: Session = Depends(get_db)):
    cur = db.get(Curriculum, curriculum_id)
    if not cur:
        raise HTTPException(status_code=404, detail="Curriculum not found")
    return curriculum_to_dict(cur)


@app.get("/api/curricula/{curriculum_id}/scheme")
def api_scheme(curriculum_id: int, db: Session = Depends(get_db)):
    cur = db.get(Curriculum, curriculum_id)
    if not cur:
        raise HTTPException(status_code=404, detail="Curriculum not found")
    if not cur.weeks:
        raise HTTPException(status_code=400, detail="This curriculum has no weeks yet")
    return {
        "html": build_scheme_html(cur),
        "title": f"Scheme of Work - {cur.subject} {cur.class_level} {cur.term}",
    }


def _generate_questions(class_level, subject, term, curriculum_block, topics_label,
                        objective_count, theory_count, difficulty, title):
    if objective_count + theory_count == 0:
        raise HTTPException(status_code=400, detail="Ask for at least one question")
    difficulty = difficulty if difficulty in ("easy", "medium", "hard", "mixed") else "mixed"
    prompt = QUESTION_PROMPT.format(
        class_level=class_level, subject=subject, term=term,
        curriculum_block=curriculum_block,
        n_obj=objective_count, n_theory=theory_count, difficulty=difficulty,
    )
    max_tokens = min(32768, 2000 + objective_count * 350 + theory_count * 500)
    raw = call_gemini_json(prompt, max_tokens=max_tokens, temperature=0.6)
    qs = _clean_questions(raw, objective_count, theory_count)
    if not qs["objective"] and not qs["theory"]:
        raise HTTPException(status_code=502, detail="AI returned no usable questions. Please try again.")
    meta = {"class_level": class_level, "subject": subject, "term": term, "topics": topics_label}
    return {
        "meta": meta,
        "questions": qs,
        "requested": {"objective": objective_count, "theory": theory_count},
        "html_student": render_questions_html(meta, qs, with_answers=False),
        "html_answers": render_questions_html(meta, qs, with_answers=True),
        "title": title,
    }


@app.post("/api/questions")
def api_questions(req: QuestionRequest, _: None = Depends(rate_limit), db: Session = Depends(get_db)):
    weeks = [db.get(Week, wid) for wid in req.week_ids]
    if any(w is None for w in weeks):
        raise HTTPException(status_code=404, detail="One or more weeks not found")
    cur = weeks[0].curriculum
    if any(w.curriculum_id != cur.id for w in weeks):
        raise HTTPException(status_code=400, detail="All weeks must belong to the same curriculum")
    weeks.sort(key=lambda w: w.week_number)

    blocks = []
    for w in weeks:
        d = week_to_dict(w)
        blocks.append(
            f"Week {d['week_number']}: {d['topic']}\n"
            f"  Objectives: {'; '.join(d['performance_objectives'])}\n"
            f"  Content: {'; '.join(d['content_knowledge'])}\n"
            f"  Evaluation guide: {'; '.join(d['evaluation_guide'])}"
        )
    return _generate_questions(
        cur.class_level, cur.subject, cur.term, "\n\n".join(blocks),
        "; ".join(f"Wk {w.week_number} {w.topic}" for w in weeks),
        req.objective_count, req.theory_count, req.difficulty,
        f"Questions - {cur.subject} {cur.class_level} Week {weeks[0].week_number}"
        + (f"-{weeks[-1].week_number}" if len(weeks) > 1 else ""),
    )


# ---------- Teacher's own topic (no uploaded curriculum needed) ----------
@app.post("/api/custom/lesson")
def api_custom_lesson(req: CustomTopicRequest, _: None = Depends(rate_limit)):
    details = [d.strip() for d in req.details.splitlines() if d.strip()]
    lesson_req = LessonNoteRequest(
        class_level=req.class_level.strip(), subject=req.subject.strip(), term="", week_number=0,
        topic=req.topic.strip(),
        performance_objectives=[], content_knowledge=details, key_competencies=[],
        learning_activities=[], teaching_resources=[], evaluation_guide=[],
    )
    html = generate_lesson_html(lesson_req)
    return {"html": html, "title": f"Lesson Note - {req.subject.strip()} {req.class_level.strip()} {req.topic.strip()}"}


@app.post("/api/custom/questions")
def api_custom_questions(req: CustomQuestionRequest, _: None = Depends(rate_limit)):
    block = f"Topic: {req.topic.strip()}"
    if req.details.strip():
        block += f"\nSub-topics / points to cover:\n{req.details.strip()}"
    return _generate_questions(
        req.class_level.strip(), req.subject.strip(), "", block, req.topic.strip(),
        req.objective_count, req.theory_count, req.difficulty,
        f"Questions - {req.subject.strip()} {req.class_level.strip()} {req.topic.strip()}",
    )


@app.post("/api/export/pdf")
def api_export_pdf(req: ExportRequest):
    pdf = html_to_pdf(req.html, req.title, req.landscape)
    return Response(content=pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{_safe_filename(req.title, "pdf")}"'})


@app.post("/api/export/docx")
def api_export_docx(req: ExportRequest):
    data = html_to_docx(sanitize_html(req.html), req.title, req.landscape)
    return Response(content=data, media_type=DOCX_MIME,
                    headers={"Content-Disposition": f'attachment; filename="{_safe_filename(req.title, "docx")}"'})


# ---------- Generate from DB ----------
def _req_from_week(w: Week, cur: Curriculum) -> LessonNoteRequest:
    d = week_to_dict(w)
    return LessonNoteRequest(
        class_level=cur.class_level,
        subject=cur.subject,
        term=cur.term,
        week_number=d["week_number"],
        topic=d["topic"],
        performance_objectives=d["performance_objectives"],
        content_knowledge=d["content_knowledge"],
        key_competencies=d["key_competencies"],
        learning_activities=d["learning_activities"],
        teaching_resources=d["teaching_resources"],
        evaluation_guide=d["evaluation_guide"],
    )


@app.post("/generate-from-db/{week_id}")
def generate_from_db(week_id: int, _: None = Depends(rate_limit), db: Session = Depends(get_db)):
    w = db.get(Week, week_id)
    if not w:
        raise HTTPException(status_code=404, detail="Week not found")
    req = _req_from_week(w, w.curriculum)
    html = generate_lesson_html(req)
    return {"html": html, "week_id": week_id, "topic": w.topic,
            "title": f"Lesson Note - {w.curriculum.subject} {w.curriculum.class_level} Week {w.week_number}"}


@app.post("/generate-from-db/{week_id}/pdf")
def generate_from_db_pdf(week_id: int, _: None = Depends(rate_limit), db: Session = Depends(get_db)):
    w = db.get(Week, week_id)
    if not w:
        raise HTTPException(status_code=404, detail="Week not found")
    req = _req_from_week(w, w.curriculum)
    html = generate_lesson_html(req)
    pdf_bytes = html_to_pdf(html, f"Lesson Note - {w.topic}")
    filename = f"Lesson_Note_{req.subject}_Week{w.week_number}.pdf".replace(" ", "_")
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@app.post("/generate-from-db/{week_id}/docx")
def generate_from_db_docx(week_id: int, _: None = Depends(rate_limit), db: Session = Depends(get_db)):
    w = db.get(Week, week_id)
    if not w:
        raise HTTPException(status_code=404, detail="Week not found")
    req = _req_from_week(w, w.curriculum)
    html = generate_lesson_html(req)
    docx_bytes = html_to_docx(html, f"Lesson Note - {w.topic}")
    filename = f"Lesson_Note_{req.subject}_Week{w.week_number}.docx".replace(" ", "_")
    return Response(
        content=docx_bytes,
        media_type=DOCX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


# ---------- Legacy endpoints (kept) ----------
@app.post("/generate-lesson-note")
def generate_lesson_note(req: LessonNoteRequest, _: None = Depends(rate_limit)):
    html = generate_lesson_html(req)
    return {"html": html, "topic": req.topic, "week": req.week_number}


@app.post("/generate-lesson-note/pdf")
def generate_lesson_pdf(req: LessonNoteRequest, _: None = Depends(rate_limit)):
    html = generate_lesson_html(req)
    pdf_bytes = html_to_pdf(html, f"Lesson Note - {req.topic}")
    filename = f"Lesson_Note_{req.subject}_Week{req.week_number}.pdf".replace(" ", "_")
    return Response(content=pdf_bytes, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.post("/generate-lesson-note/docx")
def generate_lesson_docx(req: LessonNoteRequest, _: None = Depends(rate_limit)):
    html = generate_lesson_html(req)
    docx_bytes = html_to_docx(html, f"Lesson Note - {req.topic}")
    filename = f"Lesson_Note_{req.subject}_Week{req.week_number}.docx".replace(" ", "_")
    return Response(content=docx_bytes, media_type=DOCX_MIME,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
